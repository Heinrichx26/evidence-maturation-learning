from __future__ import annotations

import ast
from collections import Counter
from pathlib import Path

import matplotlib.pyplot as plt
import pandas as pd
from matplotlib.patches import Patch
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape
from reportlab.lib.units import inch
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[2]
RESULTS = ROOT / "results"
ARTICLE = ROOT / "article" / "elsarticle_manuscript"
TABLE_DIR = ARTICLE / "tables"
FIG_DIR = ARTICLE / "figures"


STAGE_LABELS = {
    "event_aircraft_core": "Core",
    "event_aircraft_weather": "Core + weather",
    "preliminary": "Preliminary",
    "mature_factual": "Mature factual",
}

DISPLAY_LABELS = {
    "Aircraft-Fluids/misc hardware": "Aircraft fluids/hardware",
    "Environmental issues-Conditions/weather/phenomena": "Weather/phenomena",
    "Aircraft-Aircraft systems": "Aircraft systems",
    "Environmental issues-Physical environment": "Physical environment",
    "Personnel issues-Psychological": "Psychological",
    "Personnel issues-Action/decision": "Action/decision",
    "Aircraft-Aircraft oper/perf/capability": "Aircraft operation/capability",
    "Aircraft-Aircraft power plant": "Aircraft power plant",
    "Personnel issues-Task performance": "Task performance",
}


def fmt(value: object, ndigits: int = 3) -> str:
    return f"{float(value):.{ndigits}f}"


def tex_escape(text: str) -> str:
    replacements = {
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def write(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.strip() + "\n", encoding="utf-8")


def rgb(hex_color: str) -> RGBColor:
    value = hex_color.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def build_tables() -> None:
    TABLE_DIR.mkdir(parents=True, exist_ok=True)

    samples = pd.read_csv(RESULTS / "experiments" / "evidence_maturation_learning" / "samples_full.csv")
    labels = samples["labels"].map(ast.literal_eval)
    label_counts = Counter(label for row in labels for label in row)

    fold = pd.read_csv(RESULTS / "tables" / "tab_forward_validation_folds.csv")
    test_samples = int(fold["test_samples"].sum())
    summary_rows = [
        ("Observation unit", "NTSB event-aircraft record"),
        ("Study period", "2020--2025"),
        ("Training/validation units", f"{len(samples) - test_samples:,}"),
        ("Forward-test units", f"{test_samples:,}"),
        ("Final finding labels", f"{len(label_counts)} level-2 taxonomy labels"),
        ("Mean labels per unit", fmt(samples["label_count"].mean(), 2)),
        ("Mean findings per unit", fmt(samples["finding_count"].mean(), 2)),
        ("Validation design", "Test years 2023, 2024, and 2025"),
    ]
    table = [
        r"\begin{table}[t]",
        r"\centering",
        r"\caption{Data and validation summary.}",
        r"\label{tab:data_summary}",
        r"\papertablesize",
        r"\begin{tabularx}{\linewidth}{lY}",
        r"\toprule",
        r"Item & Value \\",
        r"\midrule",
    ]
    for key, value in summary_rows:
        table.append(f"{tex_escape(key)} & {tex_escape(value)} \\\\")
    table.extend(
        [
            r"\bottomrule",
            r"\end{tabularx}",
            r"\end{table}",
        ]
    )
    write(TABLE_DIR / "tab1_data_summary.tex", "\n".join(table))

    stage = pd.read_csv(RESULTS / "tables" / "tab_evidence_stage_metrics.csv")
    order = ["event_aircraft_core", "event_aircraft_weather", "preliminary", "mature_factual"]
    stage = stage.set_index("stage").loc[order].reset_index()
    best_micro = stage["micro_f1"].astype(float).max()
    best_samples = stage["samples_f1"].astype(float).max()
    best_entropy = stage["mean_binary_entropy"].astype(float).min()
    table = [
        r"\begin{table}[t]",
        r"\centering",
        r"\caption{Evidence-stage performance under forward-year validation.}",
        r"\label{tab:stage_metrics}",
        r"\papertablesize",
        r"\begin{tabularx}{\linewidth}{Yrrrrr}",
        r"\toprule",
        r"Evidence stage & Micro-F1 & Macro-F1 & Samples-F1 & Jaccard & Entropy \\",
        r"\midrule",
    ]
    for _, row in stage.iterrows():
        micro = fmt(row["micro_f1"])
        samples_f1 = fmt(row["samples_f1"])
        entropy = fmt(row["mean_binary_entropy"])
        if abs(float(row["micro_f1"]) - best_micro) < 1e-12:
            micro = rf"\textbf{{{micro}}}"
        if abs(float(row["samples_f1"]) - best_samples) < 1e-12:
            samples_f1 = rf"\textbf{{{samples_f1}}}"
        if abs(float(row["mean_binary_entropy"]) - best_entropy) < 1e-12:
            entropy = rf"\textbf{{{entropy}}}"
        table.append(
            f"{STAGE_LABELS[row['stage']]} & {micro} & {fmt(row['macro_f1'])} & "
            f"{samples_f1} & {fmt(row['jaccard_samples'])} & {entropy} \\\\"
        )
    table.extend(
        [
            r"\bottomrule",
            r"\end{tabularx}",
            r"\begin{flushleft}\footnotesize Higher F1 and Jaccard values indicate stronger closure; lower entropy indicates lower final-finding uncertainty.\end{flushleft}",
            r"\end{table}",
        ]
    )
    write(TABLE_DIR / "tab2_stage_metrics.tex", "\n".join(table))

    strong = pd.read_csv(RESULTS / "tables" / "tab_strong_baseline_preliminary_mature.csv")
    method_order = ["BR-TFIDF", "VCLDL-adapted", "SIHTC-adapted", "DyLas-adapted", "HALB-adapted"]
    strong["method"] = pd.Categorical(strong["method"], method_order, ordered=True)
    strong["stage"] = pd.Categorical(strong["stage"], ["preliminary", "mature_factual"], ordered=True)
    strong = strong.sort_values(["stage", "method"])
    best_by_stage = strong.groupby("stage", observed=True)["micro_f1"].transform("max")
    table = [
        r"\begin{table}[t]",
        r"\centering",
        r"\caption{Strong baseline comparison on preliminary and mature factual evidence.}",
        r"\label{tab:strong_baselines}",
        r"\papertablesize",
        r"\begin{tabular}{llrrrr}",
        r"\toprule",
        r"Stage & Method & Micro-F1 & Macro-F1 & Samples-F1 & Jaccard \\",
        r"\midrule",
    ]
    for idx, row in strong.iterrows():
        micro = fmt(row["micro_f1"])
        if abs(float(row["micro_f1"]) - float(best_by_stage.loc[idx])) < 1e-12:
            micro = rf"\textbf{{{micro}}}"
        table.append(
            f"{STAGE_LABELS[str(row['stage'])]} & {tex_escape(str(row['method']))} & {micro} & "
            f"{fmt(row['macro_f1'])} & {fmt(row['samples_f1'])} & {fmt(row['jaccard_samples'])} \\\\"
        )
    table.extend(
        [
            r"\bottomrule",
            r"\end{tabular}",
            r"\begin{flushleft}\footnotesize Bold values mark the strongest Micro-F1 within each evidence stage.\end{flushleft}",
            r"\end{table}",
        ]
    )
    write(TABLE_DIR / "tab3_strong_baselines.tex", "\n".join(table))

    maturation = pd.read_csv(RESULTS / "tables" / "tab_label_maturation_taxonomy.csv")
    maturation = maturation[maturation["support"].astype(float) >= 200].copy()
    maturation = maturation.sort_values("mature_minus_prelim_f1", ascending=False)
    table = [
        r"\begin{table}[t]",
        r"\centering",
        r"\caption{High-support final finding groups by maturation pattern.}",
        r"\label{tab:maturation_taxonomy}",
        r"\papertablesize",
        r"\begin{tabularx}{\linewidth}{Yrrrrl}",
        r"\toprule",
        r"Final finding group & Support & Prelim. F1 & Mature F1 & Gain & Class \\",
        r"\midrule",
    ]
    for _, row in maturation.iterrows():
        gain = fmt(row["mature_minus_prelim_f1"])
        if float(row["mature_minus_prelim_f1"]) >= 0.20:
            gain = rf"\textbf{{{gain}}}"
        label = tex_escape(DISPLAY_LABELS.get(row["label"], row["label"]))
        cls = tex_escape(row["maturation_class"].replace("_", " "))
        table.append(
            f"{label} & {int(row['support'])} & {fmt(row['preliminary_f1'])} & "
            f"{fmt(row['mature_factual_f1'])} & {gain} & {cls} \\\\"
        )
    table.extend(
        [
            r"\bottomrule",
            r"\end{tabularx}",
            r"\begin{flushleft}\footnotesize Bold gains mark finding groups with at least 0.20 mature-minus-preliminary F1 improvement.\end{flushleft}",
            r"\end{table}",
        ]
    )
    write(TABLE_DIR / "tab4_maturation_taxonomy.tex", "\n".join(table))

    acq = pd.read_csv(RESULTS / "tables" / "tab_counterfactual_evidence_acquisition.csv")
    names = {
        "weather_structured": "Weather structured fields",
        "preliminary_narrative": "Preliminary narrative",
        "mature_factual_narrative": "Mature factual narrative",
    }
    best_unc = acq["uncertainty_reduction"].astype(float).max()
    table = [
        r"\begin{table}[t]",
        r"\centering",
        r"\caption{Counterfactual evidence acquisition effects.}",
        r"\label{tab:evidence_acquisition}",
        r"\papertablesize",
        r"\begin{tabularx}{\linewidth}{Yrrrr}",
        r"\toprule",
        r"Added evidence & $\Delta$Micro-F1 & $\Delta$Macro-F1 & $\Delta$Samples-F1 & Uncertainty reduction \\",
        r"\midrule",
    ]
    for _, row in acq.iterrows():
        unc = fmt(row["uncertainty_reduction"])
        if abs(float(row["uncertainty_reduction"]) - best_unc) < 1e-12:
            unc = rf"\textbf{{{unc}}}"
        table.append(
            f"{names[row['evidence_type']]} & {fmt(row['delta_micro_f1'])} & {fmt(row['delta_macro_f1'])} & "
            f"{fmt(row['delta_samples_f1'])} & {unc} \\\\"
        )
    table.extend(
        [
            r"\bottomrule",
            r"\end{tabularx}",
            r"\begin{flushleft}\footnotesize Positive F1 deltas and positive uncertainty reduction indicate useful additional evidence.\end{flushleft}",
            r"\end{table}",
        ]
    )
    write(TABLE_DIR / "tab5_evidence_acquisition.tex", "\n".join(table))


def setup_matplotlib() -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9,
            "axes.titlesize": 10,
            "axes.labelsize": 9,
            "xtick.labelsize": 8,
            "ytick.labelsize": 8,
            "legend.fontsize": 8,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
        }
    )


def build_data_figures() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    setup_matplotlib()
    stage = pd.read_csv(RESULTS / "tables" / "tab_evidence_stage_metrics.csv")
    order = ["event_aircraft_core", "event_aircraft_weather", "preliminary", "mature_factual"]
    stage = stage.set_index("stage").loc[order].reset_index()
    x = range(len(stage))
    labels = [STAGE_LABELS[s] for s in stage["stage"]]

    fig, ax1 = plt.subplots(figsize=(7.2, 4.05))
    ax1.plot(x, stage["micro_f1"], marker="o", linewidth=2, color="#1f4e79", label="Micro-F1")
    ax1.plot(x, stage["samples_f1"], marker="s", linewidth=2, color="#5b8c5a", label="Samples-F1")
    ax1.set_ylabel("Closure score")
    ax1.set_ylim(0.30, 0.72)
    ax1.set_xticks(list(x))
    ax1.set_xticklabels(labels, rotation=12, ha="right")
    ax1.grid(axis="y", color="#d9d9d9", linewidth=0.6)
    ax2 = ax1.twinx()
    ax2.plot(x, stage["mean_binary_entropy"], marker="^", linewidth=2, color="#9a3d37", label="Entropy")
    ax2.set_ylabel("Mean binary entropy")
    ax2.set_ylim(0.50, 0.72)
    lines, names = ax1.get_legend_handles_labels()
    lines2, names2 = ax2.get_legend_handles_labels()
    ax1.legend(
        lines + lines2,
        names + names2,
        loc="upper center",
        bbox_to_anchor=(0.5, -0.25),
        ncol=3,
        frameon=False,
        handlelength=2.0,
    )
    ax1.set_title("Evidence maturation improves final finding closure")
    fig.tight_layout(rect=(0, 0.10, 1, 1))
    fig.savefig(FIG_DIR / "fig3_evidence_stage_performance.pdf", bbox_inches="tight")
    plt.close(fig)

    maturation = pd.read_csv(RESULTS / "tables" / "tab_label_maturation_taxonomy.csv")
    maturation = maturation[maturation["support"].astype(float) >= 200].copy()
    maturation = maturation.sort_values("mature_minus_prelim_f1")
    fig, ax = plt.subplots(figsize=(7.2, 4.65))
    y = range(len(maturation))
    colors_map = {
        "early_closable": "#2f6f4e",
        "late_emerging": "#9a3d37",
        "mixed": "#4d6f91",
        "unstable": "#7a6f9b",
    }
    bar_colors = [colors_map.get(c, "#666666") for c in maturation["maturation_class"]]
    values = maturation["mature_minus_prelim_f1"].astype(float)
    ax.barh(y, values, color=bar_colors, alpha=0.88)
    ax.axvline(0.0, color="#333333", linewidth=0.8)
    ax.axvline(0.20, color="#777777", linewidth=0.8, linestyle="--")
    xmax = max(0.44, float(values.max()) + 0.065)
    ax.set_xlim(0.0, xmax)
    label_offset = xmax * 0.012
    right_margin = xmax * 0.012
    for ypos, value in zip(y, values):
        x_text = float(value) + label_offset
        ha = "left"
        if x_text > xmax - right_margin:
            x_text = xmax - right_margin
            ha = "right"
        ax.text(x_text, ypos, f"{float(value):.3f}", va="center", ha=ha, fontsize=8, color="#222222")
    ax.set_yticks(list(y))
    ax.set_yticklabels([label.replace("Environmental issues-", "Env.-").replace("Personnel issues-", "Pers.-").replace("Aircraft-", "Acft.-") for label in maturation["label"]])
    ax.set_xlabel("Mature factual F1 minus preliminary F1")
    ax.set_title("Label-specific evidence maturation")
    ax.grid(axis="x", color="#d9d9d9", linewidth=0.6)
    legend_handles = [
        Patch(facecolor="#2f6f4e", alpha=0.88, label="Early closable"),
        Patch(facecolor="#9a3d37", alpha=0.88, label="Late emerging"),
        Patch(facecolor="#4d6f91", alpha=0.88, label="Mixed"),
    ]
    ax.legend(
        handles=legend_handles,
        frameon=False,
        loc="upper center",
        bbox_to_anchor=(0.5, -0.20),
        ncol=3,
        handlelength=1.4,
        columnspacing=1.6,
    )
    fig.tight_layout(rect=(0, 0.09, 1, 1))
    fig.savefig(FIG_DIR / "fig4_label_maturation_map.pdf", bbox_inches="tight")
    plt.close(fig)


def ppt_textbox(slide, left, top, width, height, text, font_size=18, bold=False, color="1F2937"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = rgb(color)
    return box


def ppt_stage(slide, left, top, width, height, title, body, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb("D1D5DB")
    frame = shape.text_frame
    frame.clear()
    p = frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = rgb("111827")
    q = frame.add_paragraph()
    q.text = body
    q.font.size = Pt(11)
    q.font.color.rgb = rgb("374151")
    return shape


def add_connector(slide, x1, y1, x2, y2, color="4B5563"):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(1.4)
    return conn


def save_pdf_workflow(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=landscape((9 * inch, 5.05 * inch)))
    w, h = landscape((9 * inch, 5.05 * inch))
    c.setFillColor(colors.HexColor("#ffffff"))
    c.rect(0, 0, w, h, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#111827"))
    c.setFont("Helvetica-Bold", 18)
    c.drawString(0.35 * inch, h - 0.45 * inch, "Preliminary-to-final evidence maturation")
    c.setFont("Helvetica", 8)
    stages = [
        ("Accident occurs", "Event and aircraft facts", "#E8F1F2"),
        ("Preliminary report", "Initial narrative and weather", "#E9F5DB"),
        ("Investigation matures", "Factual narrative expands", "#FFF4D6"),
        ("Final report", "Hierarchical findings close", "#F9E7E7"),
    ]
    y = 2.15 * inch
    box_w = 1.75 * inch
    gap = 0.38 * inch
    x = 0.35 * inch
    for title, body, fill in stages:
        c.setFillColor(colors.HexColor(fill))
        c.setStrokeColor(colors.HexColor("#CBD5E1"))
        c.roundRect(x, y, box_w, 1.25 * inch, 8, fill=1, stroke=1)
        c.setFillColor(colors.HexColor("#111827"))
        c.setFont("Helvetica-Bold", 9.5)
        c.drawString(x + 0.12 * inch, y + 0.84 * inch, title)
        c.setFont("Helvetica", 8)
        c.drawString(x + 0.12 * inch, y + 0.53 * inch, body)
        if title != "Final report":
            c.setStrokeColor(colors.HexColor("#4B5563"))
            c.line(x + box_w + 0.04 * inch, y + 0.62 * inch, x + box_w + gap - 0.08 * inch, y + 0.62 * inch)
        x += box_w + gap
    c.setFillColor(colors.HexColor("#374151"))
    c.setFont("Helvetica", 8.5)
    c.drawString(0.35 * inch, 0.65 * inch, "The gap is measured by comparing which final finding labels can be closed at each evidence stage.")
    c.save()


def save_pdf_framework(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=landscape((9 * inch, 5.05 * inch)))
    w, h = landscape((9 * inch, 5.05 * inch))
    c.setFillColor(colors.HexColor("#ffffff"))
    c.rect(0, 0, w, h, fill=1, stroke=0)
    c.setFillColor(colors.HexColor("#111827"))
    c.setFont("Helvetica-Bold", 18)
    c.drawString(0.35 * inch, h - 0.45 * inch, "Evidence-stage constrained learning framework")
    items = [
        (0.45, 3.15, "Stage-specific evidence", "Structured facts, weather, preliminary narrative, factual narrative", "#E8F1F2"),
        (3.05, 3.15, "Leakage-safe filter", "Final finding fields and trigger words are excluded", "#F1F5F9"),
        (5.65, 3.15, "Hierarchical labels", "Level-2 final finding taxonomy", "#E9F5DB"),
        (1.65, 1.45, "Forward validation", "Train on earlier years; test on later years", "#FFF4D6"),
        (4.25, 1.45, "Maturation outputs", "Early-closable, late-emerging, unstable groups", "#F9E7E7"),
    ]
    for x, y, title, body, fill in items:
        c.setFillColor(colors.HexColor(fill))
        c.setStrokeColor(colors.HexColor("#CBD5E1"))
        c.roundRect(x * inch, y * inch, 2.2 * inch, 0.9 * inch, 8, fill=1, stroke=1)
        c.setFillColor(colors.HexColor("#111827"))
        c.setFont("Helvetica-Bold", 9.5)
        c.drawString((x + 0.12) * inch, (y + 0.57) * inch, title)
        c.setFont("Helvetica", 7.5)
        c.drawString((x + 0.12) * inch, (y + 0.28) * inch, body[:58])
    c.setStrokeColor(colors.HexColor("#4B5563"))
    c.line(2.68 * inch, 3.6 * inch, 3.05 * inch, 3.6 * inch)
    c.line(5.28 * inch, 3.6 * inch, 5.65 * inch, 3.6 * inch)
    c.line(4.15 * inch, 3.15 * inch, 2.75 * inch, 2.35 * inch)
    c.line(4.95 * inch, 3.15 * inch, 5.35 * inch, 2.35 * inch)
    c.setFillColor(colors.HexColor("#374151"))
    c.setFont("Helvetica", 8.5)
    c.drawString(0.35 * inch, 0.65 * inch, "Each model is evaluated with evidence available at the corresponding investigation stage.")
    c.save()


def build_diagram_pptx() -> None:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(9)
    prs.slide_height = Inches(5.05)
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    ppt_textbox(slide, 0.35, 0.25, 8.2, 0.45, "Preliminary-to-final evidence maturation", 20, True)
    stages = [
        (0.35, "Accident occurs", "Event and aircraft facts", "E8F1F2"),
        (2.55, "Preliminary report", "Initial narrative and weather", "E9F5DB"),
        (4.75, "Investigation matures", "Factual narrative expands", "FFF4D6"),
        (6.95, "Final report", "Hierarchical findings close", "F9E7E7"),
    ]
    for left, title, body, fill in stages:
        ppt_stage(slide, left, 2.0, 1.65, 1.2, title, body, fill)
    for left in [2.0, 4.2, 6.4]:
        add_connector(slide, left, 2.6, left + 0.42, 2.6)
    ppt_textbox(
        slide,
        0.35,
        4.15,
        8.2,
        0.5,
        "The measured gap separates final findings that close from early evidence and findings that require mature investigation evidence.",
        11,
        False,
        "374151",
    )
    prs.save(FIG_DIR / "fig1_evidence_maturation_workflow.pptx")
    save_pdf_workflow(FIG_DIR / "fig1_evidence_maturation_workflow.pdf")

    prs = Presentation()
    prs.slide_width = Inches(9)
    prs.slide_height = Inches(5.05)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ppt_textbox(slide, 0.35, 0.25, 8.2, 0.45, "Evidence-stage constrained learning framework", 20, True)
    boxes = [
        (0.45, 1.1, "Stage-specific evidence", "Structured facts, weather, preliminary narrative, factual narrative", "E8F1F2"),
        (3.05, 1.1, "Leakage-safe filter", "Final finding fields and trigger words are excluded", "F1F5F9"),
        (5.65, 1.1, "Hierarchical labels", "Level-2 final finding taxonomy", "E9F5DB"),
        (1.65, 3.0, "Forward validation", "Train on earlier years; test on later years", "FFF4D6"),
        (4.25, 3.0, "Maturation outputs", "Early-closable, late-emerging, unstable groups", "F9E7E7"),
    ]
    for left, top, title, body, fill in boxes:
        ppt_stage(slide, left, top, 2.2, 0.9, title, body, fill)
    add_connector(slide, 2.65, 1.55, 3.02, 1.55)
    add_connector(slide, 5.25, 1.55, 5.62, 1.55)
    add_connector(slide, 4.0, 2.0, 2.65, 3.0)
    add_connector(slide, 4.95, 2.0, 5.35, 3.0)
    ppt_textbox(
        slide,
        0.35,
        4.35,
        8.2,
        0.35,
        "The same forward-year protocol evaluates closure at each investigation stage.",
        11,
        False,
        "374151",
    )
    prs.save(FIG_DIR / "fig2_stage_constrained_framework.pptx")
    save_pdf_framework(FIG_DIR / "fig2_stage_constrained_framework.pdf")


def main() -> None:
    build_tables()
    build_data_figures()
    build_diagram_pptx()


if __name__ == "__main__":
    main()
